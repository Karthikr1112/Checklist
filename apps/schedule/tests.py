import io
from datetime import date, time

import openpyxl
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import Role, User
from apps.schedule.models import PlanStatus, VisitPlan
from apps.stores.models import Store
from apps.visits.models import AuditType


class ScheduleExportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            email="admin@test.com",
            password="testpassword123",
            full_name="Test Admin",
            role=Role.ADMIN,
            employee_id="ADM-001",
        )
        self.client.force_login(self.user)

        self.store = Store.objects.create(
            code="STR001",
            name="Main Supermarket",
            city="Chennai",
            region="South",
        )

        self.plan = VisitPlan.objects.create(
            cluster_manager=self.user,
            store=self.store,
            planned_date=date(2026, 10, 25),
            start_time=time(9, 30),
            end_time=time(11, 30),
            audit_type=AuditType.FULL,
            status=PlanStatus.PLANNED,
            notes="Quarterly comprehensive audit",
            created_by=self.user,
        )

        # Non-planned schedule (COMPLETED) to verify exclusion
        self.completed_plan = VisitPlan.objects.create(
            cluster_manager=self.user,
            store=self.store,
            planned_date=date(2026, 10, 26),
            start_time=time(14, 0),
            end_time=time(16, 0),
            audit_type=AuditType.FULL,
            status=PlanStatus.COMPLETED,
            notes="Completed audit should not be in export",
            created_by=self.user,
        )

    def test_export_excel_endpoint_only_planned(self):
        url = reverse("schedule:export_excel")
        response = self.client.get(url, {"all": "1"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("Planned_Schedules_", response["Content-Disposition"])

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn("Schedules", wb.sheetnames)
        ws = wb["Schedules"]
        self.assertEqual(ws["A1"].value, "STORE VISIT PLANNED SCHEDULES")
        # Row 8 is first data row: Col 1 is Planned Date, Col 6 is Status, Col 8 is Store Code, Col 9 is Store Name
        self.assertEqual(ws.cell(row=8, column=6).value, "Planned")
        self.assertEqual(ws.cell(row=8, column=8).value, "STR001")
        self.assertEqual(ws.cell(row=8, column=9).value, "Main Supermarket")
        # Row 9 should be empty (completed plan excluded)
        self.assertIsNone(ws.cell(row=9, column=1).value)
        # Verify removed columns are not present in header
        self.assertEqual(ws.cell(row=7, column=26).value, "Visit Status")
        self.assertIsNone(ws.cell(row=7, column=27).value)

    def test_export_csv_endpoint_only_planned(self):
        url = reverse("schedule:export_csv")
        response = self.client.get(url, {"all": "1"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])
        self.assertIn("Planned_Schedules_", response["Content-Disposition"])

        content = response.content.decode("utf-8-sig")
        self.assertIn("Planned Date,Day of Week,Start Time", content)
        self.assertNotIn("Plan ID", content)
        self.assertNotIn(str(self.plan.id), content)
        self.assertNotIn(",Email,", content)
        self.assertNotIn(",Phone,", content)
        self.assertNotIn("Compliance %", content)
        self.assertNotIn("Total Questions", content)
        self.assertNotIn("Visit Summary", content)
        self.assertIn("STR001", content)
        self.assertIn("Main Supermarket", content)
        self.assertIn("Quarterly comprehensive audit", content)
        # Excludes non-planned plan
        self.assertNotIn("Completed audit should not be in export", content)

    def test_single_plan_export_excel(self):
        url = reverse("schedule:plan_export_excel", args=[self.plan.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb["Schedules"]
        self.assertEqual(ws.cell(row=8, column=8).value, "STR001")

    def test_single_plan_export_csv(self):
        url = reverse("schedule:plan_export_csv", args=[self.plan.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8-sig")
        self.assertIn("STR001", content)

    def test_role_restrictions(self):
        # Regular user (CLUSTER_MANAGER) cannot export
        cm_user = User.objects.create_user(
            email="cm@test.com",
            password="testpassword123",
            full_name="Test CM",
            role=Role.CLUSTER_MANAGER,
        )
        self.client.force_login(cm_user)
        response = self.client.get(reverse("schedule:export_excel"))
        self.assertEqual(response.status_code, 403)
        response = self.client.get(reverse("schedule:export_csv"))
        self.assertEqual(response.status_code, 403)

        # COO can export
        coo_user = User.objects.create_user(
            email="coo@test.com",
            password="testpassword123",
            full_name="Test COO",
            role=Role.COO,
        )
        self.client.force_login(coo_user)
        response = self.client.get(reverse("schedule:export_excel"))
        self.assertEqual(response.status_code, 200)

        # Professional user can export
        pro_user = User.objects.create_user(
            email="pro@test.com",
            password="testpassword123",
            full_name="Test Pro",
            role=Role.PROFESSIONAL_USER,
        )
        self.client.force_login(pro_user)
        response = self.client.get(reverse("schedule:export_excel"))
        self.assertEqual(response.status_code, 200)

    def test_from_to_date_filtering(self):
        url = reverse("schedule:export_csv")
        # Query date range that excludes 2026-10-25
        response = self.client.get(url, {"date_from": "2026-11-01", "date_to": "2026-11-30"})
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8-sig")
        self.assertNotIn("STR001", content)

        # Query date range that includes 2026-10-25
        response = self.client.get(url, {"date_from": "2026-10-20", "date_to": "2026-10-30"})
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8-sig")
        self.assertIn("STR001", content)

    def test_default_export_current_month(self):
        today = timezone.localdate()
        self.plan.planned_date = today
        self.plan.save()

        next_month = today.month + 1 if today.month < 12 else 1
        next_year = today.year if today.month < 12 else today.year + 1
        future_plan = VisitPlan.objects.create(
            cluster_manager=self.user,
            store=self.store,
            planned_date=date(next_year, next_month, 10),
            start_time=time(10, 0),
            end_time=time(12, 0),
            audit_type=AuditType.FULL,
            status=PlanStatus.PLANNED,
            notes="Future month planned audit",
            created_by=self.user,
        )

        url = reverse("schedule:export_csv")
        # Call without date filters -> should default to current month only
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8-sig")
        self.assertIn("Quarterly comprehensive audit", content)
        self.assertNotIn("Future month planned audit", content)

    def test_plan_start_visit_empty_coordinates_admin(self):
        from apps.checklists.models import ChecklistTemplate
        template = ChecklistTemplate.objects.create(
            name="General Visit Checklist",
            audience=ChecklistTemplate.Audience.CLUSTER_MANAGER,
            is_active=True,
            version=1,
        )
        self.plan.template = template
        self.plan.save()

        url = reverse("schedule:plan_start_visit", args=[self.plan.pk])
        # Admin posting with empty strings for user_latitude and user_longitude
        response = self.client.post(url, {"user_latitude": "", "user_longitude": ""})
        self.assertEqual(response.status_code, 302)
        self.plan.refresh_from_db()
        self.assertIsNotNone(self.plan.visit)
        self.assertIsNone(self.plan.visit.started_latitude)
        self.assertIsNone(self.plan.visit.started_longitude)

    def test_duplicate_plan_start_time_validation(self):
        url = reverse("schedule:plan_create")
        data = {
            "cluster_manager": str(self.user.pk),
            "store": str(self.store.pk),
            "planned_date": "2026-10-25",
            "start_time": "09:30",
            "end_time": "11:00",
            "audit_type": AuditType.FULL,
            "status": PlanStatus.PLANNED,
        }
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "already scheduled")

