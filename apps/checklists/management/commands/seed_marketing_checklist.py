"""Seed the Marketing Team Store Visit checklists from Marketing Teaml_Store_Visit_Checklist.xlsx.

Seeds 4 checklist templates under Audience.MARKETING_TEAM:
1. Marketing Team - Marketing Head (MH-001 .. MH-013)
2. Marketing Team - Assistant Manager (MAM-001 .. MAM-015)
3. Marketing Team - Marketing Coordinator (MC-001 .. MC-015)
4. Marketing Team - Marketing Executive (ME-001 .. ME-023)
"""
from __future__ import annotations

from pathlib import Path
from django.conf import settings
from django.core.management.base import BaseCommand
from django.db import transaction

from apps.checklists.models import ChecklistTemplate, Question, Section

BASE_DIR = Path(settings.BASE_DIR)
DEFAULT_EXCEL_PATH = BASE_DIR / "Marketing Teaml_Store_Visit_Checklist.xlsx"

# Fallback definitions in case Excel file is missing
MARKETING_TEMPLATES_DATA = {
    "Marketing Team - Marketing Head": {
        "prefix": "MH",
        "description": "Store audit & strategic review checklist for Marketing Head.",
        "items": [
            ("Strategic Review", "MH-001", "Overall store performance & footfall review", "Review MTD footfall and sales numbers", False),
            ("Campaign", "MH-002", "Campaign performance and current offers", "Review current running offers", False),
            ("Visibility", "MH-003", "Store visibility from main road", "Clear line of sight from primary approach road", True),
            ("Outdoor Branding", "MH-004", "Hoardings / outdoor branding condition", "Inspect lighting, tearing, and clean graphics", True),
            ("In-store Branding", "MH-005", "In-store branding and campaign communication", "Check entry arch, standees, overhead banners", True),
            ("Category Visibility", "MH-006", "Fashion / Home / Supermarket / Beauty / Food Court", "Review category signage prominence", False),
            ("Customer Experience", "MH-007", "Customer experience & ambience", "Check music, cleanliness, air-conditioning", False),
            ("Local Marketing", "MH-008", "Local marketing opportunities", "Catchment area partnerships, resident tie-ups", False),
            ("Digital", "MH-009", "Digital / Google Business presence and reviews", "Check recent ratings and response rate", False),
            ("Planning", "MH-010", "Upcoming festival / campaign requirements", "Identify upcoming dates & collateral needs", False),
            ("Store Feedback", "MH-011", "Store team's marketing feedback", "Record input from Store Manager & staff", False),
            ("Escalation", "MH-012", "Major gaps requiring management action", "List critical marketing bottlenecks", False),
            ("Action Plan", "MH-013", "Final action points with responsible person & deadline", "Agree on action items with SM", False),
        ],
    },
    "Marketing Team - Assistant Manager": {
        "prefix": "MAM",
        "description": "Store audit & execution review checklist for Marketing Assistant Manager.",
        "items": [
            ("Campaign Execution", "MAM-001", "Execution of approved marketing plan", "Check compliance against campaign calendar", False),
            ("Materials", "MAM-002", "All campaign materials installed", "Verify all planned touchpoints are branded", True),
            ("Offer Communication", "MAM-003", "Offer communication at customer touchpoints", "Check clarity and placement at billing/entry", True),
            ("Category Visibility", "MAM-004", "Category-wise promotional visibility", "Check department-wise promotional callouts", False),
            ("Performance", "MAM-005", "Footfall and campaign response", "Gather cashier and customer feedback on offers", False),
            ("Outdoor Branding", "MAM-006", "Outdoor branding condition", "Inspect hoardings, pole kiosks, and approach boards", True),
            ("In-store Branding", "MAM-007", "In-store branding condition", "Check for peeling, faded, or damaged prints", True),
            ("POSM", "MAM-008", "POSM availability", "Verify shelf-talkers, wobblers, danglers in place", True),
            ("Quantity", "MAM-009", "Promotional materials: received vs required", "Check dispatch quantity against floor requirements", False),
            ("Competition", "MAM-010", "Competitor promotions", "Observe and note nearby competitor offers/activities", False),
            ("Store Requirements", "MAM-011", "Store team's requirements", "Collect requests for localized banners or signage", False),
            ("Pending Activities", "MAM-012", "Pending marketing activities", "Review previously raised action items", False),
            ("Coordination", "MAM-013", "Corrections with relevant departments", "Coordinate with VM, Ops, and Facilities", False),
            ("Evidence", "MAM-014", "Photographic evidence", "Capture comprehensive store visit photographs", True),
            ("Reporting", "MAM-015", "Visit report and escalation points", "Complete and submit visit findings", False),
        ],
    },
    "Marketing Team - Marketing Coordinator": {
        "prefix": "MC",
        "description": "Store audit & collateral verification checklist for Marketing Coordinator.",
        "items": [
            ("Visit Planning", "MC-001", "Confirm visit schedule with store", "Notify Store Manager prior to arrival", False),
            ("Checklist", "MC-002", "Carry latest campaign checklist", "Verify active marketing campaign list", False),
            ("Artwork", "MC-003", "Verify approved artwork / materials", "Ensure prints match brand guidelines and approval", False),
            ("Delivery", "MC-004", "Check delivery of marketing materials", "Confirm receipt date and delivery challan", False),
            ("Quantity", "MC-005", "Quantity received vs required", "Count received items vs dispatch order", False),
            ("Execution", "MC-006", "Installation / execution status", "Inspect quality of vendor mounting and taping", True),
            ("Vendor", "MC-007", "Vendor execution quality", "Evaluate print resolution, finishing, and alignment", False),
            ("Photos", "MC-008", "Photograph all branding locations", "Take photos of all installed branding elements", True),
            ("Damage / Missing", "MC-009", "Record damaged / missing materials", "Log shortfalls or damaged shipments for reprint", True),
            ("Dates", "MC-010", "Campaign start / end dates", "Ensure expired campaign collaterals are dismantled", False),
            ("Requirements", "MC-011", "Collect store requirements", "Record signage or marketing needs from floor staff", False),
            ("Purchase", "MC-012", "Vendor / purchase-related pending items", "Verify bills and completion certificates", False),
            ("Tracker", "MC-013", "Update campaign tracker", "Record implementation status in tracking sheet", False),
            ("Issue Tracker", "MC-014", "Update store-wise issue tracker", "Flag non-compliance or vendor defects", False),
            ("Follow-up", "MC-015", "Share report and follow up until closure", "Send summary email to AM and Store Manager", False),
        ],
    },
    "Marketing Team - Marketing Executive": {
        "prefix": "ME",
        "description": "Store-level visual marketing and POS execution checklist for Marketing Executive.",
        "items": [
            ("Store Frontage", "ME-001", "Store frontage photo", "Capture wide exterior shot showing facade & sign", True),
            ("Entrance", "ME-002", "Main entrance branding", "Check entry arch, pillar wraps, and glass stickers", True),
            ("Road Visibility", "ME-003", "Road-facing visibility", "Verify facade illumination and clear visibility", True),
            ("Outdoor", "ME-004", "Hoarding / arch / lamp-post branding", "Check condition and lighting of exterior media", True),
            ("Windows", "ME-005", "Window branding", "Inspect window vinyl graphics and cleanliness", True),
            ("Standee", "ME-006", "Standee placement", "Ensure standees are placed at high-footfall spots", True),
            ("Danglers", "ME-007", "Danglers / drop-downs", "Verify ceiling danglers are aligned and undamaged", True),
            ("Category Signage", "ME-008", "Category signage", "Check overhead aisle boards and department headers", True),
            ("Offers", "ME-009", "Offer communication", "Verify tent cards, clip-ons, and shelf strips", True),
            ("Price Boards", "ME-010", "Price / offer boards", "Check price-point displays and discount callouts", True),
            ("POS", "ME-011", "Counter / POS communication", "Verify branding on cash desk and counter mats", True),
            ("Food Court", "ME-012", "Food Court branding", "Check table-top branding, menu boards, and posters", True),
            ("Supermarket", "ME-013", "Supermarket promotional branding", "Inspect end-cap headers and floor graphics", True),
            ("Fashion", "ME-014", "Fashion section branding", "Check mannequin callouts and lifestyle visuals", True),
            ("Beauty Studio", "ME-015", "Beauty Studio branding, where applicable", "Inspect cosmetic counter backdrops and mirror signs", True),
            ("Digital Screens", "ME-016", "LED / TV screen content", "Confirm active promotional videos are running on loop", True),
            ("PA", "ME-017", "PA / in-store announcement status", "Verify pre-recorded offer jingles/announcements", False),
            ("Condition", "ME-018", "Campaign material condition", "Report any torn, stained, or tilted posters", True),
            ("Competition", "ME-019", "Competitor branding nearby", "Take photo notes of neighboring retail campaigns", False),
            ("Feedback", "ME-020", "Customer / staff feedback", "Ask floor staff how customers react to the offer", False),
            ("Photos", "ME-021", "Before / after photographs", "Capture photos after fixing any misplaced signage", True),
            ("Issues", "ME-022", "Report issues immediately", "Notify store team and AM regarding missing items", False),
            ("Closure", "ME-023", "Confirm corrective action", "Verify immediate on-spot corrections were made", False),
        ],
    },
}


class Command(BaseCommand):
    help = "Seed the Marketing Team store visit checklist templates from Excel or default spec."

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            default=str(DEFAULT_EXCEL_PATH),
            help=f"Path to Marketing checklist Excel file (default: {DEFAULT_EXCEL_PATH})",
        )

    def _load_from_excel(self, excel_path: Path):
        """Attempt to read items dynamically from Excel if available."""
        try:
            import openpyxl
        except ImportError:
            return None

        if not excel_path.exists():
            return None

        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheet_mapping = {
            "Marketing Head": "Marketing Team - Marketing Head",
            "Assistant Manager": "Marketing Team - Assistant Manager",
            "Marketing Coordinator": "Marketing Team - Marketing Coordinator",
            "Marketing Executive": "Marketing Team - Marketing Executive",
        }

        extracted = {}
        for sheetname, template_name in sheet_mapping.items():
            if sheetname not in wb.sheetnames:
                continue
            ws = wb[sheetname]
            pfx = MARKETING_TEMPLATES_DATA[template_name]["prefix"]
            desc = MARKETING_TEMPLATES_DATA[template_name]["description"]
            items = []
            for r in range(8, ws.max_row + 1):
                sl = ws.cell(r, 1).value
                area = ws.cell(r, 2).value
                point = ws.cell(r, 3).value
                if sl and str(sl).strip().isdigit():
                    num = int(sl)
                    code = f"{pfx}-{num:03d}"
                    area_str = str(area or "General").strip()
                    point_str = str(point or "").strip()
                    # Check photo requirement
                    lower_point = point_str.lower()
                    req_photo = any(k in lower_point for k in ("photo", "frontage", "entrance", "window", "hoarding", "standee", "dangler", "evidence"))
                    items.append((area_str, code, point_str, "", req_photo))
            if items:
                extracted[template_name] = {
                    "prefix": pfx,
                    "description": desc,
                    "items": items,
                }
        return extracted if extracted else None

    @transaction.atomic
    def handle(self, *args, **options):
        excel_path = Path(options["file"])
        templates_data = self._load_from_excel(excel_path)
        if templates_data:
            self.stdout.write(f"Loaded checklist points from Excel: {excel_path.name}")
        else:
            templates_data = MARKETING_TEMPLATES_DATA
            self.stdout.write("Using built-in Marketing checklist specifications.")

        total_questions = 0
        for template_name, data in templates_data.items():
            template, created = ChecklistTemplate.objects.get_or_create(
                name=template_name,
                version=1,
                defaults={
                    "is_active": True,
                    "description": data["description"],
                    "audience": ChecklistTemplate.Audience.MARKETING_TEAM,
                },
            )
            if not created and template.audience != ChecklistTemplate.Audience.MARKETING_TEAM:
                template.audience = ChecklistTemplate.Audience.MARKETING_TEAM
                template.save(update_fields=["audience", "updated_at"])

            action_text = "Created" if created else "Updated"
            self.stdout.write(self.style.SUCCESS(f"{action_text} template: {template_name}"))

            # Track section orders
            section_cache: dict[str, Section] = {}
            section_order: dict[str, int] = {}
            for item in data["items"]:
                sec_name = item[0]
                if sec_name not in section_order:
                    section_order[sec_name] = len(section_order) + 1

            for q_order, (sec_name, code, text, evidence, requires_photo) in enumerate(data["items"], start=1):
                section = section_cache.get(sec_name)
                if section is None:
                    section, _ = Section.objects.get_or_create(
                        template=template,
                        name=sec_name,
                        defaults={"order": section_order[sec_name]},
                    )
                    section_cache[sec_name] = section

                Question.objects.update_or_create(
                    section=section,
                    code=code,
                    defaults={
                        "text": text,
                        "evidence_hint": evidence,
                        "requires_photo": requires_photo,
                        "order": q_order,
                    },
                )
                total_questions += 1

        self.stdout.write(
            self.style.SUCCESS(
                f"\nSuccessfully seeded {len(templates_data)} Marketing Team templates with {total_questions} total questions!"
            )
        )
